# -*- coding: utf-8 -*-
"""从《玄武江湖人物成长和天性与伤害计算.xlsx》导出面板攻击计算器用数据。"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(r"D:\QQ\data\玄武江湖人物成长和天性与伤害计算.xlsx")
FALLBACK_XLSX = ROOT / "game-data" / "玄武江湖人物成长和天性与伤害计算.xlsx"
OUT_JS = ROOT / "js" / "panel-attack-data.js"
COMPENDIUM_JS = ROOT / "js" / "equipment-compendium-data.js"
MARTIAL_COMPENDIUM_JS = ROOT / "js" / "martial-compendium-data.js"

BASE_ATTACK_CONST = 1569  # 与 Excel B4 公式一致
ATTR_ATTACK_FLAT = 65536  # 装备图鉴「攻击 N」固定
ATTR_ATTACK_PCT = 131072  # 装备图鉴「攻击 N%」
ATTR_CRIT_DMG_PCT = 16777216  # 装备图鉴「暴伤 N%」（副属性等）
WEAPON_EQUIP_TYPES = frozenset({1, 2, 3})  # 拳 / 剑 / 刀
GEAR_SLOT_TYPES = {5: "armors", 6: "rings", 7: "bracers"}  # 甲 / 戒 / 腕


def xlsx_path() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    if DEFAULT_XLSX.is_file():
        return DEFAULT_XLSX
    if FALLBACK_XLSX.is_file():
        return FALLBACK_XLSX
    raise SystemExit(f"未找到 xlsx: {DEFAULT_XLSX} 或 {FALLBACK_XLSX}")


def num(v) -> float | int | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def clean_float(v: float | int | None, places: int = 4) -> float | int | None:
    """去掉 Excel / IEEE 浮点噪声（如 1.0999999999999999 → 1.1）。"""
    if v is None:
        return None
    if isinstance(v, int) and not isinstance(v, bool):
        return v
    r = round(float(v), places)
    if r == int(r):
        return int(r)
    return r


def load_char_growth(ws) -> list[dict]:
    rows = []
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, 3).value
        if not name or str(name).strip() in ("", "全角色"):
            continue
        atk = num(ws.cell(r, 5).value)
        hp = num(ws.cell(r, 4).value)
        if atk is None and hp is None:
            continue
        rows.append(
            {
                "name": str(name).strip(),
                "weaponType": str(ws.cell(r, 2).value or "").strip(),
                "hpGrowth": clean_float(hp),
                "atkGrowth": clean_float(atk),
            }
        )
    return rows


def load_compendium_items() -> list[dict]:
    if not COMPENDIUM_JS.is_file():
        return []
    raw = COMPENDIUM_JS.read_text(encoding="utf-8")
    start = raw.find("[")
    end = raw.rfind("]")
    if start < 0 or end <= start:
        return []
    return json.loads(raw[start : end + 1])


def attack_stats_from_item(it: dict) -> tuple[int, int]:
    flat = 0
    pct = 0
    for key in ("mainAttr", "subAttr"):
        attr = it.get(key)
        if not attr:
            continue
        code = attr.get("code")
        v = attr.get("value")
        if v is None:
            continue
        if code == ATTR_ATTACK_FLAT:
            flat = max(flat, int(v))
        elif code == ATTR_ATTACK_PCT:
            pct = max(pct, int(v))
    for attr in it.get("enhanceAttrs") or []:
        code = attr.get("code")
        v = attr.get("value")
        if v is None:
            continue
        if code == ATTR_ATTACK_FLAT:
            flat = max(flat, int(v))
        elif code == ATTR_ATTACK_PCT:
            pct = max(pct, int(v))
    return flat, pct


def weapon_main_sub_line(attr: dict | None) -> tuple[int, int, int]:
    """单条主/副属性：攻击固定、攻击%、暴伤%（面板用小数，此处仍用整数百分比）。"""
    if not attr:
        return 0, 0, 0
    code = attr.get("code")
    v = attr.get("value")
    if v is None:
        return 0, 0, 0
    iv = int(v)
    if code == ATTR_ATTACK_FLAT:
        return iv, 0, 0
    if code == ATTR_ATTACK_PCT:
        return 0, iv, 0
    if code == ATTR_CRIT_DMG_PCT:
        return 0, 0, iv
    return 0, 0, 0


def load_compendium_weapon_stats() -> dict[str, dict]:
    """武器图鉴主属性 + 副属性（同名多品质取各项最高，不含强化词条）。"""
    by_name: dict[str, dict] = {}
    for it in load_compendium_items():
        if it.get("type") not in WEAPON_EQUIP_TYPES:
            continue
        name = str(it.get("name") or "").strip()
        if not name:
            continue
        slot = by_name.setdefault(
            name, {"attack": 0, "attackPct": 0, "critDmgPct": 0}
        )
        for key in ("mainAttr", "subAttr"):
            flat, pct, crit = weapon_main_sub_line(it.get(key))
            slot["attack"] = max(slot["attack"], flat)
            slot["attackPct"] = max(slot["attackPct"], pct)
            slot["critDmgPct"] = max(slot["critDmgPct"], crit)
    return by_name


def load_compendium_weapon_forge_options() -> dict[str, list[dict]]:
    """可锻造属性 = 图鉴 enhanceAttrs；同名武器按属性 code 取最高 value。"""
    by_name: dict[str, dict[int, dict]] = {}
    for it in load_compendium_items():
        if it.get("type") not in WEAPON_EQUIP_TYPES:
            continue
        name = str(it.get("name") or "").strip()
        if not name:
            continue
        bucket = by_name.setdefault(name, {})
        for attr in it.get("enhanceAttrs") or []:
            code = int(attr.get("code") or 0)
            if not code:
                continue
            v = attr.get("value")
            if v is None:
                continue
            iv = int(v)
            prev = bucket.get(code)
            if not prev or iv > int(prev.get("value") or 0):
                bucket[code] = {
                    "code": code,
                    "text": str(attr.get("text") or "").strip(),
                    "value": iv,
                }
    return {
        name: sorted(codes.values(), key=lambda x: (x["text"], x["code"]))
        for name, codes in by_name.items()
        if codes
    }


def load_compendium_attack_gear() -> dict[str, list[dict]]:
    """衣服 / 戒指 / 护腕：图鉴全量（同名取攻击固定、攻击%最高，无攻项为 0）。"""
    buckets: dict[int, dict[str, dict]] = {t: {} for t in GEAR_SLOT_TYPES}
    for it in load_compendium_items():
        t = it.get("type")
        if t not in buckets:
            continue
        name = str(it.get("name") or "").strip()
        if not name:
            continue
        flat, pct = attack_stats_from_item(it)
        slot = buckets[t]
        if name not in slot:
            slot[name] = {"name": name, "attackFlat": 0, "attackPct": 0}
        slot[name]["attackFlat"] = max(slot[name]["attackFlat"], flat)
        slot[name]["attackPct"] = max(slot[name]["attackPct"], pct)
    out: dict[str, list[dict]] = {}
    for t, key in GEAR_SLOT_TYPES.items():
        rows = sorted(buckets[t].values(), key=lambda x: x["name"])
        out[key] = rows
    return out


def compendium_weapon_names() -> set[str]:
    names: set[str] = set()
    for it in load_compendium_items():
        if it.get("type") not in WEAPON_EQUIP_TYPES:
            continue
        name = str(it.get("name") or "").strip()
        if name:
            names.add(name)
    return names


def load_weapons_from_compendium(
    comp_by_name: dict[str, dict], forge_by_name: dict[str, list[dict]]
) -> list[dict]:
    """武器列表与数值均来自装备图鉴（主/副属性），可锻造词条见 forgeOptions。"""
    rows = []
    for name_s in sorted(compendium_weapon_names()):
        comp = comp_by_name.get(name_s, {})
        attack = int(comp.get("attack") or 0)
        attack_pct = int(comp.get("attackPct") or 0)
        crit_pct = int(comp.get("critDmgPct") or 0)
        row: dict = {
            "name": name_s,
            "attack": attack,
            "attackPct": attack_pct,
        }
        if crit_pct > 0:
            row["critDmg"] = clean_float(crit_pct / 100)
        forge_opts = forge_by_name.get(name_s)
        if forge_opts:
            row["forgeOptions"] = forge_opts
        rows.append(row)
    return rows


def load_internal(ws) -> list[dict]:
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        atk_pct = num(ws.cell(r, 5).value)  # E列 攻击
        crit_dmg = num(ws.cell(r, 9).value)  # I列 爆伤
        if atk_pct is None and crit_dmg is None:
            continue
        rows.append(
            {
                "name": str(name).strip(),
                "grade": str(ws.cell(r, 2).value or "").strip(),
                "attackPct": clean_float(atk_pct),
                "critDmgPct": clean_float(crit_dmg),
            }
        )
    return rows


def load_martial_compendium_crit() -> dict[str, float]:
    """武功图鉴 baojiEffect = 游戏内武功爆伤，优先于社区 xlsx I 列。"""
    if not MARTIAL_COMPENDIUM_JS.is_file():
        return {}
    raw = MARTIAL_COMPENDIUM_JS.read_text(encoding="utf-8")
    start = raw.find("[")
    end = raw.rfind("]")
    if start < 0 or end <= start:
        return {}
    items = json.loads(raw[start : end + 1])
    out: dict[str, float] = {}
    for it in items:
        name = str(it.get("name") or "").strip()
        eff = it.get("baojiEffect")
        if name and eff is not None:
            out[name] = float(eff)
    return out


def load_martial(ws) -> list[dict]:
    crit_from_comp = load_martial_compendium_crit()
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name or str(name).strip().endswith("…"):
            continue
        power = num(ws.cell(r, 4).value)  # D 威力
        speed = num(ws.cell(r, 5).value)
        coeff = num(ws.cell(r, 8).value)  # H 系数
        crit_dmg = num(ws.cell(r, 9).value)  # I 爆伤
        if power is None:
            continue
        skill_name = str(name).strip()
        crit = crit_from_comp.get(skill_name, crit_dmg)
        rows.append(
            {
                "name": skill_name,
                "grade": str(ws.cell(r, 2).value or "").strip(),
                "power": int(power) if power == int(power) else power,
                "speed": clean_float(speed),
                "coeff": clean_float(coeff),
                "critDmg": clean_float(crit),
            }
        )
    return rows


def main() -> None:
    path = xlsx_path()
    wb = openpyxl.load_workbook(path, data_only=True)

    gear = load_compendium_attack_gear()
    payload = {
        "source": path.name,
        "baseAttackConst": BASE_ATTACK_CONST,
        "characters": load_char_growth(wb["人物成长与天性"]),
        "martial": load_martial(wb["武学"]),
        "internal": load_internal(wb["内功"]),
        "weapons": load_weapons_from_compendium(
            load_compendium_weapon_stats(),
            load_compendium_weapon_forge_options(),
        ),
        "armors": gear["armors"],
        "rings": gear["rings"],
        "bracers": gear["bracers"],
        "gearSource": "js/equipment-compendium-data.js",
        "formulaNotes": [
            "基础攻击 B4 = round(1569 × 攻击成长)",
            "核心 B13 = floor(武功威力 + 基础攻击 × (100 + 攻击丹) / 100)",
            "任督内功 = floor(基础攻击 × 攻击内功任督 × 0.1)",
            "武功小任督 = floor(武功威力 × 武功小任督% / 100)",
            "内功%段 C13 = round(B13 × 内功攻击小数)",
            "其他武功% = 武功威力 ×「其他武功加成%」/ 100（与表 F4 一致）",
            "装备 E11 = 武器 + 衣服 + 戒指 + 护腕（固定 + B13×%）",
            "技艺 F11 = round(B13 × 技艺加成比例)，比例见页面技艺区",
            "面板 A13 = floor(B13 + 任督内功 + 小任督 + C13 + 其他武功% + E11 + F11)",
            "大任督 A4 = A13 + floor(A13 × 武功大任督 × 5%)",
            "对单总伤 A6 ≈ round(A4 × 总爆伤 × 武功系数 + 附伤)，上限约 ×1.1",
        ],
    }

    OUT_JS.write_text(
        "// 由 scripts/build-panel-attack-data.py 生成，请勿手改。\n"
        "window.PANEL_ATTACK_DATA = "
        + json.dumps(payload, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    print(
        f"角色 {len(payload['characters'])} · 武学 {len(payload['martial'])} · 内功 {len(payload['internal'])} · "
        f"武器 {len(payload['weapons'])} · 衣 {len(payload['armors'])} · 戒 {len(payload['rings'])} · 腕 {len(payload['bracers'])}"
    )
    print(f"已写入: {OUT_JS}")


if __name__ == "__main__":
    main()
